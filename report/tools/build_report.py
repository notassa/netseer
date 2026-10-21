"""Build the revised FYP2 report (.docx) from report/src/*.txt.

Usage:  python report/tools/build_report.py
Then:   powershell -File report/tools/finalize.ps1   (updates TOC/lists in Word and exports PDF)

Missing screenshots (docs/screenshots/*.png) are rendered as placeholder boxes; drop the
PNGs in and re-run both commands to embed them.
"""
import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "src"
OUT = ROOT / "FYP2-Report-revised.docx"
TEXT_W = 8.27 - 1.5 - 1.0  # A4 width minus left/right margins (inches)

TITLE = "NetSeer: AI-Assisted Alert Interpretation System for Intrusion Detection Monitoring in Small Networks"
NAME = "MUHAMMAD DANISH IRFAN BIN SELJALI"
SID = "52215224164"
MONTH = "OCTOBER 2026"


# ---------- low-level helpers ----------
ORDER = {
    "tcPr": ["cnfStyle", "tcW", "gridSpan", "hMerge", "vMerge", "tcBorders", "shd", "noWrap", "tcMar",
             "textDirection", "tcFitText", "vAlign", "hideMark"],
    "tblPr": ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize",
              "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook"],
    "sectPr": ["headerReference", "footerReference", "footnotePr", "endnotePr", "type", "pgSz", "pgMar",
               "paperSrc", "pgBorders", "lnNumType", "pgNumType", "cols", "formProt", "vAlign", "noEndnote",
               "titlePg", "textDirection", "bidi", "rtlGutter", "docGrid", "printerSettings"],
}


def insert_ordered(parent, child):
    """Insert child into parent respecting the OOXML schema sequence."""
    order = ORDER[parent.tag.split("}")[1]]
    name = child.tag.split("}")[1]
    idx = order.index(name)
    for i, existing in enumerate(parent):
        ename = existing.tag.split("}")[1]
        if ename in order and order.index(ename) > idx:
            parent.insert(i, child)
            return child
    parent.append(child)
    return child

def set_cell_shading(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    insert_ordered(tcPr, shd)


def set_cell_text_direction(cell, direction):
    tcPr = cell._tc.get_or_add_tcPr()
    td = OxmlElement("w:textDirection")
    td.set(qn("w:val"), direction)
    insert_ordered(tcPr, td)


def set_cell_valign(cell, val):
    tcPr = cell._tc.get_or_add_tcPr()
    v = OxmlElement("w:vAlign")
    v.set(qn("w:val"), val)
    insert_ordered(tcPr, v)


def set_table_borders(table, size=4, val="single"):
    tblPr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), val)
        el.set(qn("w:sz"), str(size))
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "000000" if val != "nil" else "auto")
        borders.append(el)
    insert_ordered(tblPr, borders)


def set_cell_border(cell, edge, val="single", size=8):
    tcPr = cell._tc.get_or_add_tcPr()
    b = tcPr.find(qn("w:tcBorders"))
    if b is None:
        b = OxmlElement("w:tcBorders")
        insert_ordered(tcPr, b)
    el = OxmlElement(f"w:{edge}")
    el.set(qn("w:val"), val)
    el.set(qn("w:sz"), str(size))
    el.set(qn("w:space"), "0")
    el.set(qn("w:color"), "000000")
    b.append(el)


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    h = OxmlElement("w:tblHeader")
    h.set(qn("w:val"), "true")
    trPr.append(h)


def no_split(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:cantSplit"))


def add_field(paragraph, instr, placeholder="Right-click and choose Update Field."):
    def fld(t):
        r = paragraph.add_run()
        e = OxmlElement("w:fldChar")
        e.set(qn("w:fldCharType"), t)
        r._r.append(e)

    fld("begin")
    r = paragraph.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = instr
    r._r.append(it)
    fld("separate")
    paragraph.add_run(placeholder)
    fld("end")


def clear_theme_fonts(style):
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rFonts.get(qn(a)) is not None:
            del rFonts.attrib[qn(a)]
    for a in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rFonts.set(qn(a), "Arial")


def para_fmt(p, align=None, before=None, after=None, spacing=None, keep_next=None):
    pf = p.paragraph_format
    if align is not None:
        pf.alignment = align
    if before is not None:
        pf.space_before = Pt(before)
    if after is not None:
        pf.space_after = Pt(after)
    if spacing is not None:
        pf.line_spacing = spacing
    if keep_next is not None:
        pf.keep_with_next = keep_next
    return p


INLINE = re.compile(r"(\*\*.+?\*\*|\[\[.+?\]\])")


def add_inline(p, text, size=None, bold=False, font=None):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = p.add_run(part[2:-2])
            r.bold = True
        elif part.startswith("[[") and part.endswith("]]"):
            r = p.add_run(part[2:-2])
            r.font.highlight_color = WD_COLOR_INDEX.YELLOW
            r.bold = bold
        else:
            r = p.add_run(part)
            r.bold = bold
        if size:
            r.font.size = Pt(size)
        if font:
            r.font.name = font
    return p


# ---------- document set-up ----------
def setup(doc):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
    sec.left_margin, sec.right_margin = Inches(1.5), Inches(1.0)
    sec.top_margin, sec.bottom_margin = Inches(1.0), Inches(1.0)

    st = doc.styles
    normal = st["Normal"]
    normal.font.name, normal.font.size = "Arial", Pt(12)
    clear_theme_fonts(normal)
    pf = normal.paragraph_format
    pf.line_spacing, pf.space_after, pf.space_before = 1.5, Pt(8), Pt(0)
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for name, size, align, before, after in (
        ("Heading 1", 14, WD_ALIGN_PARAGRAPH.CENTER, 0, 24),
        ("Heading 2", 12, WD_ALIGN_PARAGRAPH.LEFT, 14, 6),
        ("Heading 3", 12, WD_ALIGN_PARAGRAPH.LEFT, 10, 6),
    ):
        h = st[name]
        h.font.name, h.font.size, h.font.bold, h.font.italic = "Arial", Pt(size), True, False
        h.font.color.rgb = RGBColor(0, 0, 0)
        clear_theme_fonts(h)
        h.paragraph_format.alignment = align
        h.paragraph_format.space_before, h.paragraph_format.space_after = Pt(before), Pt(after)
        h.paragraph_format.line_spacing = 1.5
        h.paragraph_format.keep_with_next = True

    from docx.enum.style import WD_STYLE_TYPE

    for name, align in (("Table Caption", WD_ALIGN_PARAGRAPH.LEFT), ("Figure Caption", WD_ALIGN_PARAGRAPH.CENTER)):
        s = st.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        s.base_style = normal
        s.font.bold = True
        s.paragraph_format.alignment = align
        s.paragraph_format.space_before, s.paragraph_format.space_after = Pt(6), Pt(6)
        s.paragraph_format.keep_with_next = name == "Table Caption"

    s = st.add_style("Table Text", WD_STYLE_TYPE.PARAGRAPH)
    s.base_style = normal
    s.font.size = Pt(10.5)
    s.paragraph_format.line_spacing = 1.0
    s.paragraph_format.space_after = Pt(2)
    s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    for lvl_style in ("List Bullet",):
        ls = st[lvl_style]
        ls.font.name, ls.font.size = "Arial", Pt(12)
        ls.paragraph_format.line_spacing = 1.5
        ls.paragraph_format.space_after = Pt(4)
        ls.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY



def page_break(doc):
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def heading1(doc, text, new_page=True):
    if new_page:
        page_break(doc)
    return doc.add_heading(text, level=1)


# ---------- front matter ----------
def centered(doc, text, bold=False, size=12, before=0, after=0, spacing=1.15):
    p = doc.add_paragraph()
    para_fmt(p, WD_ALIGN_PARAGRAPH.CENTER, before, after, spacing)
    r = p.add_run(text)
    r.bold, r.font.size = bold, Pt(size)
    return p


def cover(doc):
    t = doc.add_table(rows=1, cols=2)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    set_table_borders(t, val="nil")
    row = t.rows[0]
    row.height = Inches(9.4)
    row.height_rule = WD_ROW_HEIGHT_RULE.EXACTLY
    spine, body = row.cells
    spine.width, body.width = Inches(0.6), Inches(TEXT_W - 0.6)
    t.columns[0].width, t.columns[1].width = Inches(0.6), Inches(TEXT_W - 0.6)
    set_cell_text_direction(spine, "tbRl")
    set_cell_border(spine, "right", size=6)
    sp = spine.paragraphs[0]
    para_fmt(sp, WD_ALIGN_PARAGRAPH.CENTER, 0, 0, 1.0)
    r = sp.add_run("BIT (HONS) IN COMPUTER SYSTEM SECURITY" + " " * 40 + "2026")
    r.bold, r.font.size = True, Pt(11)
    set_cell_valign(spine, "center")

    body.paragraphs[0].text = ""
    def bp(text, bold=True, before=0, after=0):
        p = body.add_paragraph()
        para_fmt(p, WD_ALIGN_PARAGRAPH.CENTER, before, after, 1.15)
        rr = p.add_run(text)
        rr.bold, rr.font.size = bold, Pt(12)

    bp(TITLE, before=60)
    bp(NAME, before=190)
    bp(SID)
    bp("UNIVERSITI KUALA LUMPUR", before=200)
    bp(MONTH, before=40)
    page_break(doc)

    centered(doc, TITLE, bold=True, before=40)
    centered(doc, NAME, bold=True, before=120)
    centered(doc, SID, bold=True)
    centered(doc, "Report Submitted to Fulfill the Partial Requirements", before=150)
    centered(doc, "For the Bachelor of Information Technology (Hons)")
    centered(doc, "in Computer System Security")
    centered(doc, "Universiti Kuala Lumpur")
    centered(doc, MONTH, bold=True, before=150)


def signature_block(doc, left_lines, right_lines, image=None):
    t = doc.add_table(rows=1, cols=2)
    set_table_borders(t, val="nil")
    l, r = t.rows[0].cells
    l.width, r.width = Inches(2.6), Inches(TEXT_W - 2.6)
    t.columns[0].width, t.columns[1].width = Inches(2.6), Inches(TEXT_W - 2.6)
    l.paragraphs[0].text = ""
    for ln in left_lines:
        p = l.add_paragraph(ln)
        para_fmt(p, WD_ALIGN_PARAGRAPH.LEFT, 0, 0, 1.15)
    r.paragraphs[0].text = ""
    if image and Path(image).exists():
        p = r.add_paragraph()
        para_fmt(p, WD_ALIGN_PARAGRAPH.LEFT, 0, 0, 1.0)
        p.add_run().add_picture(str(image), width=Inches(0.8))
    for ln in right_lines:
        p = r.add_paragraph(ln)
        para_fmt(p, WD_ALIGN_PARAGRAPH.LEFT, 0, 0, 1.15)


def front_pages(doc):
    # numbered section starts at DECLARATION, page 1
    sec = doc.add_section(WD_SECTION.NEW_PAGE)
    sec.footer.is_linked_to_previous = False
    pgNumType = OxmlElement("w:pgNumType")
    pgNumType.set(qn("w:start"), "1")
    insert_ordered(sec._sectPr, pgNumType)
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    add_field(fp, "PAGE", "1")
    doc.sections[0].footer.is_linked_to_previous = False

    heading1(doc, "DECLARATION", new_page=False)
    p = centered(doc, "I declare that this report is my original work and all references have been cited adequately as required by the University", before=170, after=170, spacing=1.5)
    signature_block(doc, ["", "", "Date: ...................."],
                    ["Signature: ..................................", f"Full Name: {NAME}", f"ID Number: {SID}"],
                    image=ROOT / "assets" / "signature.png")

    heading1(doc, "APPROVAL PAGE")
    centered(doc, "I certify that I have supervised and read this report and that, in my opinion, it conforms to acceptable standards of scholarly presentation and fulfils the partial requirements for the Bachelor of Information Technology (Hons) in Computer System Security.", before=170, after=170, spacing=1.5)
    signature_block(doc, ["", "Date: ...................."],
                    ["Signature: ..................................", "Supervisor: Dr. Norliana Muslim"])

    heading1(doc, "DEDICATION")
    centered(doc, "This project is dedicated to my beloved family, whose constant prayers, patience and encouragement made this work possible, and to the many small organisations and everyday network administrators whose security deserves to be made understandable.", before=60, spacing=1.5)

    heading1(doc, "ACKNOWLEDGEMENT")
    for t in (
        "Assalamualaikum W.B.T.",
        "First and foremost, I would like to thank Allah S.W.T., for with His blessings I was able to complete this Final Year Project successfully.",
        "My sincere gratitude goes to my supervisor, Dr. Norliana Muslim, who guided, supported and advised me throughout the duration of this project.",
        "I would also like to thank my lecturers and friends, who gave me inspiration, support and motivation during the completion of this work.",
        "Finally, I am deeply grateful to my family for their prayers and continuous support. Thank you to everyone who was directly or indirectly involved in helping me complete this project. Special thanks to Oriental Tech, who helped and reviewed the project.",
    ):
        doc.add_paragraph(t)

    heading1(doc, "TABLE OF CONTENTS")
    add_field(doc.add_paragraph(), 'TOC \\o "1-3" \\h \\z \\u')
    heading1(doc, "LIST OF TABLES")
    add_field(doc.add_paragraph(), 'TOC \\h \\z \\t "Table Caption,1"')
    heading1(doc, "LIST OF FIGURES")
    add_field(doc.add_paragraph(), 'TOC \\h \\z \\t "Figure Caption,1"')

    heading1(doc, "LIST OF ABBREVIATIONS")
    abbr = [
        ("ACP", "Alert-Context Processing"), ("AI", "Artificial Intelligence"),
        ("API", "Application Programming Interface"), ("ASVS", "Application Security Verification Standard"),
        ("ATT&CK", "Adversarial Tactics, Techniques and Common Knowledge"), ("CSRF", "Cross-Site Request Forgery"),
        ("CVSS", "Common Vulnerability Scoring System"), ("DHCP", "Dynamic Host Configuration Protocol"),
        ("DNS", "Domain Name System"), ("FTS", "Full-Text Search"),
        ("HIDS", "Host-based Intrusion Detection System"), ("IDS", "Intrusion Detection System"),
        ("IPS", "Intrusion Prevention System"), ("JSON", "JavaScript Object Notation"),
        ("LAN", "Local Area Network"), ("LLM", "Large Language Model"),
        ("MAC", "Media Access Control"), ("NIDS", "Network-based Intrusion Detection System"),
        ("NIST", "National Institute of Standards and Technology"), ("NSM", "Network Security Monitoring"),
        ("OUI", "Organisationally Unique Identifier"), ("OWASP", "Open Worldwide Application Security Project"),
        ("PCAP", "Packet Capture"), ("RAD", "Rapid Application Development"),
        ("RAG", "Retrieval-Augmented Generation"), ("RBAC", "Role-Based Access Control"),
        ("SIEM", "Security Information and Event Management"), ("SME", "Small and Medium-sized Enterprise"),
        ("SNI", "Server Name Indication"), ("SOHO", "Small Office / Home Office"),
        ("SPAN", "Switched Port Analyzer (port mirroring)"), ("SUS", "System Usability Scale"),
        ("TLS", "Transport Layer Security"), ("UI", "User Interface"),
        ("XAI", "Explainable Artificial Intelligence"),
    ]
    t = doc.add_table(rows=0, cols=2)
    set_table_borders(t, val="nil")
    t.columns[0].width, t.columns[1].width = Inches(1.6), Inches(TEXT_W - 1.6)
    for a, m in [("Abbreviation", "Meaning")] + abbr:
        cells = t.add_row().cells
        cells[0].width, cells[1].width = Inches(1.6), Inches(TEXT_W - 1.6)
        for c, txt in zip(cells, (a, m)):
            p = c.paragraphs[0]
            para_fmt(p, WD_ALIGN_PARAGRAPH.LEFT, 0, 1, 1.0)
            r = p.add_run(txt)
            r.bold = a == "Abbreviation"


# ---------- body parsing ----------
def add_table(doc, caption, widths, rows):
    doc.add_paragraph(caption, style="Table Caption")
    ncol = len(rows[0])
    widths = widths or [100 / ncol] * ncol
    t = doc.add_table(rows=0, cols=ncol)
    t.style = doc.styles["Table Grid"]
    t.autofit = False
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    keep = len(rows) <= 9
    for i, row in enumerate(rows):
        cells = t.add_row().cells
        no_split(t.rows[-1])
        if i == 0:
            repeat_header(t.rows[-1])
        for j, c in enumerate(cells):
            c.width = Inches(TEXT_W * widths[j] / 100)
            p = c.paragraphs[0]
            p.style = doc.styles["Table Text"]
            p.paragraph_format.keep_with_next = i == 0 or (keep and i < len(rows) - 1)
            add_inline(p, row[j] if j < len(row) else "", bold=(i == 0), size=9 if ncol >= 6 else None)
            if i == 0:
                set_cell_shading(c, "E7E6E6")
    for j, col in enumerate(t.columns):
        col.width = Inches(TEXT_W * widths[j] / 100)
    spacer = doc.add_paragraph()
    para_fmt(spacer, after=0, spacing=1.0)


def add_matrix(doc, caption, rows):
    doc.add_paragraph(caption, style="Table Caption")
    ncol = len(rows[0])
    first = 1.55
    other = (TEXT_W - first) / (ncol - 1)
    t = doc.add_table(rows=0, cols=ncol)
    t.style = doc.styles["Table Grid"]
    t.autofit = False
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        cells = t.add_row().cells
        if i == 0:
            t.rows[-1].height = Inches(2.1)
            t.rows[-1].height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
        for j, c in enumerate(cells):
            c.width = Inches(first if j == 0 else other)
            p = c.paragraphs[0]
            p.style = doc.styles["Table Text"]
            txt = row[j]
            if i > 0 and j > 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = p.add_run("✓" if txt == "Y" else "✗")
                r.font.name = "Segoe UI Symbol"
                r.font.size = Pt(11)
                set_cell_valign(c, "center")
            else:
                add_inline(p, txt, bold=(i == 0))
            p.paragraph_format.keep_with_next = i < len(rows) - 1
            if i == 0:
                set_cell_shading(c, "E7E6E6")
                if j > 0:
                    set_cell_text_direction(c, "btLr")
    for j, col in enumerate(t.columns):
        col.width = Inches(first if j == 0 else other)
    para_fmt(doc.add_paragraph(), after=0, spacing=1.0)


def add_figure(doc, path, caption, width):
    img = (ROOT / path).resolve()
    if img.exists():
        p = doc.add_paragraph()
        para_fmt(p, WD_ALIGN_PARAGRAPH.CENTER, 6, 0, 1.0, keep_next=True)
        p.add_run().add_picture(str(img), width=Inches(min(width, TEXT_W)))
    else:
        t = doc.add_table(rows=1, cols=1)
        t.style = doc.styles["Table Grid"]
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        c = t.rows[0].cells[0]
        c.width = Inches(TEXT_W - 0.4)
        t.rows[0].height = Inches(2.2)
        set_cell_valign(c, "center")
        set_cell_shading(c, "F2F2F2")
        p = c.paragraphs[0]
        para_fmt(p, WD_ALIGN_PARAGRAPH.CENTER, 0, 0, 1.0)
        r = p.add_run(f"[Insert screenshot: {Path(path).name}]")
        r.font.highlight_color = WD_COLOR_INDEX.YELLOW
    doc.add_paragraph(caption, style="Figure Caption")


def parse_body(doc, text):
    lines = text.splitlines()
    i = 0
    in_refs = False
    while i < len(lines):
        ln = lines[i].rstrip()
        i += 1
        if not ln.strip():
            continue
        if in_refs and not ln.startswith("@"):
            p = doc.add_paragraph()
            para_fmt(p, WD_ALIGN_PARAGRAPH.LEFT, 0, 8, 1.5)
            p.paragraph_format.left_indent = Inches(0.5)
            p.paragraph_format.first_line_indent = Inches(-0.5)
            p.add_run(ln)
            continue
        if ln.startswith("@front ") or ln.startswith("@chapter "):
            heading1(doc, ln.split(" ", 1)[1])
        elif ln.startswith("@h2 "):
            doc.add_heading(ln[4:], level=2)
        elif ln.startswith("@h3 "):
            doc.add_heading(ln[4:], level=3)
        elif ln.startswith("@table ") or ln.startswith("@matrix "):
            is_matrix = ln.startswith("@matrix ")
            head = ln.split(" ", 1)[1]
            caption, _, w = head.partition(" | ")
            widths = [float(x) for x in w.split(",")] if w.strip() else None
            rows = []
            end = "@endmatrix" if is_matrix else "@endtable"
            while not lines[i].startswith(end):
                rows.append([c.strip() for c in lines[i].split(" | ")])
                i += 1
            i += 1
            if is_matrix:
                add_matrix(doc, caption.strip(), rows)
            else:
                add_table(doc, caption.strip(), widths, rows)
        elif ln.startswith("@figure "):
            path, caption, width = [x.strip() for x in ln[8:].split(" | ")]
            add_figure(doc, path, caption, float(width))
        elif ln.startswith("@source "):
            p = doc.add_paragraph()
            para_fmt(p, WD_ALIGN_PARAGRAPH.CENTER, 0, 8, 1.0)
            r = p.add_run(ln[8:])
            r.italic, r.font.size = True, Pt(9)
        elif ln.startswith("@eq "):
            p = doc.add_paragraph()
            para_fmt(p, WD_ALIGN_PARAGRAPH.CENTER, 2, 6, 1.5)
            r = p.add_run(ln[4:])
            r.italic = True
        elif ln == "@bullets":
            while lines[i] != "@endbullets":
                p = doc.add_paragraph(style="List Bullet")
                add_inline(p, lines[i][2:])
                i += 1
            i += 1
        elif ln == "@list":
            while lines[i] != "@endlist":
                num, _, rest = lines[i].partition(" ")
                p = doc.add_paragraph()
                p.paragraph_format.left_indent = Inches(0.5)
                p.paragraph_format.first_line_indent = Inches(-0.4)
                add_inline(p, f"{num}\t{rest}")
                i += 1
            i += 1
        elif ln == "@refs":
            in_refs = True
        else:
            add_inline(doc.add_paragraph(), ln)


def main():
    doc = Document()
    setup(doc)
    cover(doc)
    front_pages(doc)
    for f in sorted(SRC.glob("*.txt")):
        parse_body(doc, f.read_text(encoding="utf-8"))
    doc.save(OUT)
    print(f"wrote {OUT}")


if __name__ == "__main__":
    main()
